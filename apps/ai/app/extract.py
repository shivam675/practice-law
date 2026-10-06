"""Document text extraction.

This is the only part of the platform that parses attacker-supplied binary
formats. PDF and DOCX parsers are a long-running source of CVEs, so the work is
isolated here: a container with no network route out, a read-only filesystem,
a memory cap and hard limits on how much of a document will be read.

Extraction returns text and, where the format can report one honestly, a page
count. DOCX has no page count without rendering, so it returns ``None`` rather
than a guess: a fabricated page count would produce confident wrong findings
in a structural check on a graded submission.
"""

from __future__ import annotations

import io
import logging
import zipfile
from dataclasses import dataclass

log = logging.getLogger(__name__)

# Bounds on the work a single request can cause. A document beyond these is
# not a memorial; it is either a mistake or an attempt to exhaust the service.
MAX_BYTES = 50 * 1024 * 1024
MAX_PAGES = 300
MAX_CHARS = 4_000_000


class UnsupportedFormat(Exception):
    """The bytes do not match a format this service can read."""


class ExtractionFailed(Exception):
    """The format was recognised but the content could not be read."""


@dataclass
class Extraction:
    format: str
    text: str
    pages: int | None
    truncated: bool

    def as_dict(self) -> dict:
        return {
            "format": self.format,
            "text": self.text,
            "pages": self.pages,
            "chars": len(self.text),
            "truncated": self.truncated,
        }


def sniff(data: bytes) -> str:
    """Identify a format from its leading bytes.

    The declared filename and content type come from the client and are not
    evidence of anything.
    """
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        # Both .docx and .zip start this way; the part names decide.
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                names = set(zf.namelist())
        except zipfile.BadZipFile as exc:
            raise UnsupportedFormat("not a readable zip container") from exc
        if "word/document.xml" in names:
            return "docx"
        raise UnsupportedFormat("zip archive is not a Word document")
    if _looks_like_text(data):
        return "txt"
    raise UnsupportedFormat("unrecognised file format")


def _looks_like_text(data: bytes) -> bool:
    sample = data[:4096]
    if b"\x00" in sample:
        return False
    try:
        sample.decode("utf-8")
    except UnicodeDecodeError:
        return False
    return True


def extract(data: bytes, declared_format: str | None = None) -> Extraction:
    if len(data) > MAX_BYTES:
        raise ExtractionFailed(f"document exceeds {MAX_BYTES} bytes")
    if not data:
        raise UnsupportedFormat("empty document")

    fmt = sniff(data)
    if declared_format and declared_format != fmt:
        # Worth recording: a mismatch is either a renamed file or an attempt to
        # slip one parser's input past another's validation.
        log.warning("declared format %s does not match sniffed format %s", declared_format, fmt)

    if fmt == "pdf":
        return _extract_pdf(data)
    if fmt == "docx":
        return _extract_docx(data)
    return _extract_text(data)


def _truncate(text: str) -> tuple[str, bool]:
    if len(text) <= MAX_CHARS:
        return text, False
    return text[:MAX_CHARS], True


def _extract_pdf(data: bytes) -> Extraction:
    import pdfplumber

    parts: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            page_count = len(pdf.pages)
            for index, page in enumerate(pdf.pages):
                if index >= MAX_PAGES:
                    break
                parts.append(page.extract_text() or "")
    except Exception as exc:  # pdfminer raises a wide variety of types
        raise ExtractionFailed(f"could not read PDF: {exc}") from exc

    text, truncated = _truncate("\n".join(parts))
    return Extraction(
        format="pdf",
        text=text,
        pages=page_count,
        truncated=truncated or page_count > MAX_PAGES,
    )


def _extract_docx(data: bytes) -> Extraction:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionFailed(f"could not read Word document: {exc}") from exc

    lines: list[str] = [p.text for p in document.paragraphs]

    # Index of authorities and tables of contents are often laid out as tables,
    # and a structural check that cannot see them reports them missing.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            joined = "  ".join(c for c in cells if c)
            if joined:
                lines.append(joined)

    text, truncated = _truncate("\n".join(lines))
    # A page count would have to be rendered, so none is reported.
    return Extraction(format="docx", text=text, pages=None, truncated=truncated)


def _extract_text(data: bytes) -> Extraction:
    try:
        decoded = data.decode("utf-8")
    except UnicodeDecodeError:
        decoded = data.decode("utf-8", errors="replace")
    text, truncated = _truncate(decoded)
    return Extraction(format="txt", text=text, pages=None, truncated=truncated)
