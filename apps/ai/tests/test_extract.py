"""Extraction tests.

These run without a model, a network or a database, so they stay fast enough
to run on every change.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from app.extract import (
    ExtractionFailed,
    UnsupportedFormat,
    extract,
    sniff,
)


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_sniff_identifies_formats():
    assert sniff(b"%PDF-1.7\nrest") == "pdf"
    assert sniff(make_docx(["hello"])) == "docx"
    assert sniff(b"STATEMENT OF FACTS\nplain text") == "txt"


def test_sniff_rejects_a_plain_zip():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("notes.txt", "not a memorial")

    with pytest.raises(UnsupportedFormat):
        sniff(buffer.getvalue())


def test_sniff_rejects_binary_junk():
    with pytest.raises(UnsupportedFormat):
        sniff(b"\x00\x01\x02\x03binary")


def test_empty_document_is_refused():
    with pytest.raises(UnsupportedFormat):
        extract(b"")


def test_docx_extraction_includes_paragraphs():
    data = make_docx(["STATEMENT OF FACTS", "On 6 March the order issued."])
    result = extract(data)

    assert result.format == "docx"
    assert "STATEMENT OF FACTS" in result.text
    assert "6 March" in result.text
    # A page count cannot be known without rendering, so none is claimed.
    assert result.pages is None


def test_docx_extraction_includes_tables():
    # An index of authorities is usually a table. A structural check that
    # cannot see tables reports the section missing.
    data = make_docx(
        ["INDEX OF AUTHORITIES"],
        table=[["Anuradha Bhasin v. Union of India", "(2020) 3 SCC 637"]],
    )
    result = extract(data)

    assert "Anuradha Bhasin" in result.text
    assert "(2020) 3 SCC 637" in result.text


def test_corrupt_docx_fails_cleanly():
    # A zip that declares the right part but holds garbage.
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as zf:
        zf.writestr("word/document.xml", "not xml at all <<<")

    with pytest.raises(ExtractionFailed):
        extract(buffer.getvalue())


def test_corrupt_pdf_fails_cleanly():
    with pytest.raises(ExtractionFailed):
        extract(b"%PDF-1.7\nthis is not a pdf body")


def test_text_extraction_round_trip():
    result = extract("PRAYER\nWherefore it is prayed.".encode())

    assert result.format == "txt"
    assert "Wherefore" in result.text
    assert result.truncated is False


def test_oversize_document_is_refused():
    from app.extract import MAX_BYTES

    with pytest.raises(ExtractionFailed):
        extract(b"%PDF-" + b"x" * (MAX_BYTES + 1))


def test_long_text_is_truncated_and_says_so():
    from app.extract import MAX_CHARS

    result = extract(("word " * (MAX_CHARS // 2)).encode())

    assert result.truncated is True
    assert len(result.text) == MAX_CHARS
